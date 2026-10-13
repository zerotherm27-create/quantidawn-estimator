import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill


class ExcelGenerator:
    @staticmethod
    def create_estimation_sheet(output_path, project_name, region, parsed_data, counted_items):
        """Generates the estimator audit workbook.

        parsed_data is the dict returned by LLMParser.parse_scope_and_materials:
        {"scope_summary": str, "materials": [...], "flags": [...]}.
        This is the fix for the original spec, where the AI's text analysis
        was shown in the UI but never actually reached the exported file -
        every row below the vision-detected symbols was hardcoded placeholder
        text regardless of what the document actually said.
        """
        parsed_data = parsed_data or {}
        materials = parsed_data.get("materials", [])
        flags = parsed_data.get("flags", [])
        scope_summary = parsed_data.get("scope_summary", "")

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "AI Estimate"

        header_font = Font(name="Arial", size=12, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        section_font = Font(name="Arial", size=10, bold=True, italic=True)
        summary_font = Font(name="Arial", size=10, italic=True)
        flag_font = Font(name="Arial", size=10, color="9C5700")
        align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)

        # Project header
        ws.append([f"Project: {project_name}", "", f"Region: {region}"])
        ws["A1"].font = Font(bold=True, size=11)
        row_num = 2

        if scope_summary:
            ws.append([f"Scope summary: {scope_summary}"])
            ws.cell(row=row_num, column=1).font = summary_font
            ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=7)
            ws.cell(row=row_num, column=1).alignment = align_left
            row_num += 1

        row_num += 1  # blank spacer

        # Column headers
        headers = ["Item ID", "Description", "AI Qty / Count", "Unit", "Unit Rate", "Total Cost", "Verified?"]
        header_row = row_num
        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=header_row, column=col_idx, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = align_left
        row_num += 1

        # Section 1: vision-detected symbols
        if counted_items:
            ws.cell(row=row_num, column=1, value="-- DETECTED VIA VISUAL TAKEOFF (VERIFY COUNTS) --").font = section_font
            row_num += 1
            for i, (name, count) in enumerate(counted_items.items(), start=1):
                ws.append([f"V-{i}", f"Detected symbol: {name}", count, "EA", 0.00, f"=C{row_num}*E{row_num}", "PENDING"])
                row_num += 1
            row_num += 1

        # Section 2: AI-extracted scope and materials - this is the part that
        # was previously hardcoded and discarded. Now it's the real output.
        ws.cell(row=row_num, column=1, value="-- EXTRACTED FROM DOCUMENT TEXT (SET RATE AND VERIFY) --").font = section_font
        row_num += 1

        if materials:
            for i, item in enumerate(materials, start=1):
                description = item.get("description", "Unnamed item")
                category = item.get("category", "")
                label = f"[{category}] {description}" if category else description
                quantity = item.get("quantity")
                unit = item.get("unit", "")
                notes = item.get("notes", "")
                if notes:
                    label = f"{label} - {notes}"
                ws.append([f"T-{i}", label, quantity, unit, 0.00, f"=C{row_num}*E{row_num}", "PENDING"])
                row_num += 1
        else:
            ws.cell(row=row_num, column=2, value="No materials extracted - check the source document has readable scope text.")
            row_num += 1

        # Flags section - mirrors the QA-flags concept from the live web demo
        if flags:
            row_num += 1
            ws.cell(row=row_num, column=1, value="FLAGS - REVIEW BEFORE FINALIZING").font = Font(bold=True, color="9C5700")
            row_num += 1
            for f in flags:
                ws.cell(row=row_num, column=1, value=f"- {f}").font = flag_font
                ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=7)
                row_num += 1

        # Column widths
        widths = {"A": 10, "B": 55, "C": 14, "D": 10, "E": 12, "F": 14, "G": 12}
        for col_letter, width in widths.items():
            ws.column_dimensions[col_letter].width = width

        wb.save(output_path)
