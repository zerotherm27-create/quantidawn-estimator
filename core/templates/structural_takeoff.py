"""Writes the Westgarth/Boxshall-style TAKEOFF sheet: groups, live formulas, SUMMARY rows."""
import re

import openpyxl
from openpyxl.styles import Font, PatternFill

from core.takeoff_builder import WESTGARTH, group_order

HEAD = ["", "Description", "Units", "Diameter (m.)", "Length (m.)", "Width /thickness(m.)",
        "Height (m.)", "Area (m2)", "Count (each)", "Volume (m3)", "Status"]
SUM_COL = {"m3": "J", "m2": "H", "tons": "I", "m": "E", "each": "I"}


def _row_formulas(l, r):
    """Return {column: value-or-formula} for one line, formulas wherever dims exist."""
    c = {"A": l.mark, "B": l.description, "C": l.unit}
    for col, v in (("E", l.length), ("F", l.width), ("G", l.height), ("I", l.count)):
        if v is not None and not (l.unit == "tons" and col == "I"):
            c[col] = round(v, 4)
    mult = f"*I{r}" if l.count and l.unit != "tons" else ""
    if l.unit == "m3":
        if l.length and l.width:
            c["J"] = f"=E{r}*F{r}*G{r}{mult}"
        else:
            c["H"] = round(l.area, 4)
            c["J"] = f"=H{r}*G{r}"
    elif l.unit == "m2":
        if l.length and l.width and l.group != "FORMWORKS":
            c["H"] = f"=E{r}*F{r}{mult}"
        elif l.group == "FORMWORKS" and l.width:              # hob: both faces
            c["H"] = f"=E{r}*2*G{r}"
        elif l.group == "FORMWORKS":
            c["H"] = f"=E{r}*G{r}"
        else:
            c["H"] = round(l.area, 4)
    elif l.unit == "m":
        pass
    elif l.unit == "tons":
        kg = re.search(r"([\d.]+)kg/m2", l.description)
        c["H"] = round(l.area, 4)
        c["I"] = f"=H{r}*{kg.group(1)}/1000" if kg else round(l.count, 4)
    return c


def _residential_row(l, r):
    """One residential line with a live formula in the unit's own column; the factor and openings sit inside the formula."""
    c = {"A": l.mark, "B": l.description, "C": l.unit}
    L, W, H, N, A = l.length, l.width, l.height, l.count, l.area
    R = f"*{round(l.factor, 6)}" if l.factor not in (None, 1) else ""
    out = {"m2": "H", "m3": "J", "m": "E", "each": "I", "tons": "I"}[l.unit]
    if l.calc == "plan" and A is not None and L is None:
        c["H"] = round(A, 4)
    elif l.calc == "plan":                                                  # width x height x count (windows, doors)
        c["E"], c["F"], c["I"] = round(L, 4), round(W, 4), N
        c["H"] = f"=E{r}*F{r}*I{r}"
    elif l.calc == "prism":
        c["H"], c["G"] = round(A, 4), round(H, 4)
        c["J"] = f"=H{r}*G{r}"
    elif l.calc in ("edge", "wall"):
        c["E"], c["G"] = round(L, 4), round(H, 4)
        o = f"-{round(l.openings, 4)}" if l.calc == "wall" and l.openings else ""
        c[out] = f"=(E{r}*G{r}{o}){R}" if o else f"=E{r}*G{r}{R}"
    elif l.calc == "len":
        c["E"] = round(L, 4)
        c[out] = f"=E{r}{R}" if R else round(L, 4)
    elif l.calc == "each":
        c["I"] = N
        if H is not None:
            c["G"] = round(H, 4)
            c["E"] = f"=I{r}*G{r}"
    elif l.calc == "bar":                                                   # total bar length x kg per m
        c["E"] = round(A, 4)
        c["I"] = f"=E{r}*{round(l.k, 4)}/1000"
    return c


def write_structural_takeoff(lines, path, skipped=(), profile=WESTGARTH, sheet="TAKEOFF", pricing=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    priced = bool(pricing and pricing.get("on"))
    head = HEAD + (["Rate", "Material cost", "Labour hrs", "Labour rate", "Labour cost"] if priced else [])
    group_cost = []          # (group, summary row) for the pricing sheet
    for i, h in enumerate(head, 1):
        cell = ws.cell(row=1, column=i, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", start_color="1F4E78")
    r = 3
    present = []
    for l in lines:
        if l.group not in present:
            present.append(l.group)
    for group in group_order(profile) + [g for g in present if g not in group_order(profile)]:
        glines = [l for l in lines if l.group == group]
        if not glines:
            continue
        ws.cell(row=r, column=1, value=group).font = Font(bold=True)
        r += 2
        first = r
        residential = profile.trade == "residential"
        for l in glines:
            for col, v in (_residential_row if residential else _row_formulas)(l, r).items():
                ws[f"{col}{r}"] = v
            ws[f"K{r}"] = getattr(l, "status", "PENDING")
            if priced:
                qc = SUM_COL.get(l.unit, "J")
                if getattr(l, "rate", None) is not None:
                    ws[f"L{r}"] = l.rate
                    ws[f"M{r}"] = f"={qc}{r}*L{r}"
                if getattr(l, "labour_hours", None) is not None:      # hours = quantity x hours per unit
                    ws[f"N{r}"] = f"={qc}{r}*{round(l.labour_hours, 4)}"
                    if getattr(l, "labour_rate", None) is not None:
                        ws[f"O{r}"] = l.labour_rate
                        ws[f"P{r}"] = f"=N{r}*O{r}"
            r += 1
        if residential:                       # a trade mixes m2, m, each and tonnes: one SUMIF row per unit
            last = r - 1
            for unit in dict.fromkeys(l.unit for l in glines):
                col = SUM_COL[unit]
                ws[f"B{r}"] = f"SUMMARY ({unit})"
                ws[f"{col}{r}"] = f'=SUMIF(C{first}:C{last},"{unit}",{col}{first}:{col}{last})'
                ws[f"B{r}"].font = ws[f"{col}{r}"].font = Font(bold=True)
                r += 1
            r -= 1
        else:
            col = SUM_COL.get(glines[0].unit, "J")
            label_col = {"J": "I", "H": "G", "I": "F", "E": "D"}[col]
            ws[f"{label_col}{r}"] = "SUMMARY"
            ws[f"{col}{r}"] = f"=SUM({col}{first}:{col}{r - 1})"
            ws[f"{col}{r}"].font = Font(bold=True)
        if priced:
            ws[f"M{r}"] = f"=SUM(M{first}:M{r - 1})"
            ws[f"P{r}"] = f"=SUM(P{first}:P{r - 1})"
            ws[f"M{r}"].font = ws[f"P{r}"].font = Font(bold=True)
            group_cost.append((group, r))
        r += 2
    if skipped:
        ws.cell(row=r, column=1, value="REVIEW - NOT TAKEN OFF (missing dimensions or no rule)").font = Font(bold=True, color="9C5700")
        r += 1
        for label, why in skipped:
            ws[f"B{r}"] = label
            ws[f"C{r}"] = why
            r += 1
    for col, w in zip("ABCDEFGHIJKLMNOP", (14, 70, 8, 12, 12, 14, 12, 12, 12, 12, 11, 12, 14, 12, 12, 14)):
        ws.column_dimensions[col].width = w
    if priced:
        pct = pricing.get("pct", {})
        ps = wb.create_sheet("Pricing summary")
        ps["A1"] = "Pricing summary"; ps["A1"].font = Font(bold=True, size=13)
        ps["A2"] = f"Currency: {pricing.get('currency', 'AUD')}. Lines with no rate or labour are not included."
        for col, t in (("B", "Materials"), ("C", "Labour"), ("D", "Total")):
            ps[f"{col}3"] = t; ps[f"{col}3"].font = Font(bold=True)
        row = first = 4
        for g, rr in group_cost:
            ps[f"A{row}"] = g
            ps[f"B{row}"] = f"='{sheet}'!M{rr}"
            ps[f"C{row}"] = f"='{sheet}'!P{rr}"
            ps[f"D{row}"] = f"=B{row}+C{row}"
            row += 1
        ps[f"A{row}"] = "Subtotal"
        for col in "BCD":
            ps[f"{col}{row}"] = f"=SUM({col}{first}:{col}{row - 1})"
        sub = row; row += 1
        for label, key, default in (("Contingency", "contingency", 10), ("Margin / overheads and profit", "margin", 15)):
            ps[f"A{row}"] = f"{label} ({pct.get(key, default)}%)"
            ps[f"D{row}"] = f"=D{sub}*{pct.get(key, default)}/100"
            row += 1
        ps[f"A{row}"] = "Subtotal (excl. GST)"; ps[f"D{row}"] = f"=SUM(D{sub}:D{row - 1})"; ex = row; row += 1
        ps[f"A{row}"] = f"GST ({pct.get('gst', 10)}%)"; ps[f"D{row}"] = f"=D{ex}*{pct.get('gst', 10)}/100"; row += 1
        ps[f"A{row}"] = "Total (incl. GST)"; ps[f"D{row}"] = f"=D{ex}+D{row - 1}"
        for rr in range(first, row + 1):
            for col in "BCD":
                ps[f"{col}{rr}"].number_format = "#,##0.00"
        for rr in (sub, ex, row):
            for col in "ABCD":
                ps[f"{col}{rr}"].font = Font(bold=True)
        ps.column_dimensions["A"].width = 48
        for col in "BCD":
            ps.column_dimensions[col].width = 16
    wb.save(path)
