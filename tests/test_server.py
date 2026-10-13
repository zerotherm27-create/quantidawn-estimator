import io, json
import openpyxl
from server import app
from core.project import sample_project

c = app.test_client()
print("projects:", [p["id"] for p in c.get("/api/projects").json])

r = c.get("/api/project/boxshall"); p = r.json
print("boxshall:", r.status_code, len(p["lines"]), "lines,", len(p["pages"]), "pages, groups", len(p["groups"]))
print("page image:", c.get(p["pages"][0]["img"]).status_code)
print("demo ai:", sorted({l["ag"] for l in c.get("/api/project/boxshall?demo_ai=1").json["lines"]}))

# state round trip
c.put("/api/project/boxshall/state", json={"v": 2, "x": 1}); print("state:", c.get("/api/project/boxshall").json["state"]["x"]); c.delete("/api/project/boxshall/state")

# export edited lines
r = c.post("/api/project/boxshall/export", json={"lines": p["lines"]})
wb = openpyxl.load_workbook(io.BytesIO(r.data)); ws = wb["TAKEOFF"]
print("export:", r.status_code, "rows", ws.max_row, "status col:", {ws[f'K{i}'].value for i in range(1, ws.max_row + 1)} - {None})

# client template: Westgarth blank template + Westgarth lines
w = sample_project("westgarth")
with open("samples/82-westgarth/template_blank.xlsx", "rb") as f:
    r = c.post("/api/project/westgarth/template", data={"template": (f, "client.xlsx"), "lines": json.dumps(w["lines"])}, content_type="multipart/form-data")
t = r.json
print("template:", r.status_code, t["sheets"][0]["kind"], "typed rows", t["sheets"][0]["typed"], "| matched", len(t["matches"]), "unmatched lines", t["unmatched_lines"], "unfilled rows", t["unfilled_rows"])
r = c.post(f"/api/project/westgarth/template/{t['id']}/fill", json={"lines": w["lines"]})
print("fill:", r.status_code, r.headers["X-Fill-Report"])
filled = openpyxl.load_workbook(io.BytesIO(r.data))["TAKEOFF"]
print("filled cells with numbers in dimension cols:", sum(1 for row in filled.iter_rows(min_col=5, max_col=9) for x in row if isinstance(x.value, (int, float))))

# upload a small marked-up pdf
with open("/Users/jojo/Downloads/BC26140-Turnout-Rooms-Bangholm-Training-Camp-Marked-Up.pdf", "rb") as f:
    r = c.post("/api/projects", data={"files": (f, "bangholm.pdf")}, content_type="multipart/form-data")
print("upload:", r.status_code, r.json)
if r.status_code == 200:
    up = c.get(f"/api/project/{r.json['id']}").json
    print("uploaded:", len(up["pages"]), "pages, paper", {x["paper"] for x in up["pages"]}, "scale", up["pages"][0]["scale"]["status"], "lines", len(up["lines"]))
    print("export with no lines:", c.post(f"/api/project/{r.json['id']}/export", json={"lines": []}).status_code)

# marked-up PDF: measured shapes drawn onto the sheets; the response is a link, because the file is the whole drawing set
import fitz
shape = {"kind": "area", "page": 0, "pts": [[100, 100], [300, 100], [300, 250]], "label": "PF1", "text": "9.00 m2"}
r = c.post("/api/project/boxshall/markup-pdf", json={"shapes": [shape, {"kind": "nope"}]})
f = c.get(r.json["url"])
print("markup pdf:", r.status_code, r.json["name"], "| file", f.status_code, f.mimetype, "pages", len(fitz.open(stream=f.data, filetype="pdf")),
      "| no shapes:", c.post("/api/project/boxshall/markup-pdf", json={"shapes": []}).status_code)

# project owner: recorded on upload, listed, and absent on samples
with open("samples/boxshall-brighton/marked_up.pdf", "rb") as f:
    own = c.post("/api/projects", data={"files": (f, "ownercheck.pdf")}, content_type="multipart/form-data").json["id"]
rows = {x["id"]: x for x in c.get("/api/projects").json}
print("owner:", rows[own]["owner"], "| sample has owner:", "owner" in rows["boxshall"])
