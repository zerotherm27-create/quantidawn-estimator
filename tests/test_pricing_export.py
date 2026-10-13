import io, json
import openpyxl
from server import app
from core.project import sample_project

c = app.test_client()
p = sample_project("boxshall")
lines = p["lines"]
with open("samples/pricing/EXAMPLE_price_list_not_real_rates.csv", "rb") as f:
    r = c.post("/api/project/boxshall/prices", data={"prices": (f, "example.csv"), "lines": json.dumps(lines)}, content_type="multipart/form-data")
j = r.json
print("price upload:", r.status_code, j["entries"], "entries", {s: sum(m["status"] == s for m in j["matches"]) for s in ("matched", "suggested", "none")})
rates = {m["id"]: m["rate"] for m in j["matches"] if m["status"] == "matched"}
for l in lines:
    if l["id"] in rates:
        l["rate"] = rates[l["id"]]
r = c.post("/api/project/boxshall/export", json={"lines": lines, "pricing": {"on": True, "pct": {"contingency": 10, "margin": 15, "gst": 10}, "currency": "AUD"}})
wb = openpyxl.load_workbook(io.BytesIO(r.data))
ws, ps = wb["TAKEOFF"], wb["Pricing summary"]
print("export:", r.status_code, "| sheets", wb.sheetnames, "| rate cells", sum(1 for row in ws.iter_rows(min_col=12, max_col=12) for x in row if isinstance(x.value, (int, float))))
print("summary rows:", [(ps[f"A{i}"].value, ps[f"B{i}"].value) for i in range(4, ps.max_row + 1)][-6:])
off = c.post("/api/project/boxshall/export", json={"lines": lines})
print("pricing off -> no rate columns / summary sheet:", openpyxl.load_workbook(io.BytesIO(off.data)).sheetnames, "| bad file:", c.post("/api/project/boxshall/prices", data={"prices": (io.BytesIO(b"x"), "a.txt"), "lines": "[]"}, content_type="multipart/form-data").status_code)
