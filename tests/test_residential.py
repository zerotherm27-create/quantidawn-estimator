"""Residential takeoff rules: quantities, formulas, profile choice, export. No API calls: the legend is the stored AI reading."""
import json
import os
import tempfile

from openpyxl import load_workbook

from core.calc import n_of, qty
from core.takeoff_builder import PROFILES, build_takeoff
from core.templates.structural_takeoff import write_structural_takeoff

fails = 0
def check(cond, msg):
    global fails
    print(("PASS " if cond else "FAIL ") + msg); fails += (not cond)

RES = PROFILES["residential"]
def E(label, q, unit):
    return {"label": label, "quantity": q, "unit": unit}
def build(*entries):
    return build_takeoff(list(entries), RES)
def one(lines, mark, unit=None):
    hit = [l for l in lines if l.mark == mark and (unit is None or l.unit == unit)]
    assert len(hit) == 1, (mark, [l.mark for l in lines])
    return hit[0]
def near(a, b, tol=0.01):
    return a is not None and abs(a - b) <= max(tol, 1e-6 * abs(b))

# hand-calculated from the Clinton BOQ
lines, sk = build(E("100mm Slab Supply and Pour on top of waffle Pod/EPS Expanded Polystyrence/Vapor Barrier -  U1 (SL82 Mesh)-138.14m2", 1.0, "EA"))
check(not sk and {l.mark for l in lines} == {"SLAB", "MESH", "EPS", "VAPOUR BARRIER"}, "a slab row gives pour, concrete volume, mesh, EPS and vapour barrier")
check(near(one(lines, "SLAB", "m3").quantity(), 13.814), "slab concrete volume = area x thickness (13.814 m3)")
check(near(one(lines, "MESH").quantity(), 138.14), "mesh = slab area")

lines, _ = build(E("Ground Slab Reo - N12 X 4-U2", 377.6, "M"))
check(near(one(lines, "SLAB REO").quantity(), 377.6 * 4 * 0.888 / 1000, 0.0005), "slab reo tonnes = length x 4 bars x 0.888 kg/m")
check(near(one(lines, "RIBS").quantity(), 75.52), "ribs and edge beams = length x 0.2 m3/m")

lines, _ = build(E("Termite Protection-U1", 50.8, "M"))
check(near(one(lines, "TERMITE").quantity(), 50.8) and near(one(lines, "EDGE FORM").quantity(), 19.558), "termite length and its 385 mm edge formwork (19.558 m2)")

lines, _ = build(E("Slab Edge Formworks - 100mm THICK-U1", 64.1, "M"), E("50mm Slab Edge Formwork - U1", 8.7, "M"))
check(sorted(round(l.quantity(), 3) for l in lines) == [0.435, 6.41], "edge formwork area = length x thickness from the label (100mm and 50mm)")

lines, _ = build(E("230Lx110Wmmx76H Brickworks-2550mm H (Less 34.35m2 Openings)-U2", 41.9, "M"))
br = one(lines, "BRICKWORK")
check(near(br.quantity(), 41.9 * 2.55 - 34.35) and br.openings == 34.35, "brickwork area = length x height less the openings in the label")
check(near(one(lines, "BRICKS").quantity(), (41.9 * 2.55 - 34.35) * 400 / 7, 0.5), "bricks follow the same area x 57.14 per m2")

lines, _ = build(E("190mm Blockworks - 2550mm H-U2", 6.0, "M"))
check(near(one(lines, "BLOCKWORK").quantity(), 15.3) and near(one(lines, "BLOCKS").quantity(), 191.25) and near(one(lines, "CORE FILL").quantity(), 0.612),
      "blockwork gives area, blocks and core fill")

lines, _ = build(E("Aluminum Sliding Window/with Mesh - 1800-1800 U2", 2.0, "EA"))
check(near(one(lines, "WINDOW").quantity(), 6.48) and near(one(lines, "BLIND").quantity(), 6.48), "two 1800 x 1800 windows = 6.48 m2, and the same area of blinds")
lines, _ = build(E("Aluminum Sliding Door/Insect-screen doors-2100-4200", 1.0, "EA"))
check(near(one(lines, "SLIDING DOOR").quantity(), 8.82), "a 2100 x 4200 sliding door = 8.82 m2")
lines, _ = build(E("Solid Core Door (entry) - 870 x 2100 U1", 1.0, "EA"), E("Panel Lift Door - Auto - 2100-4800 U1", 1.0, "EA"))
check(near(one(lines, "DOOR").quantity(), 1.827) and near(one(lines, "GARAGE DOOR").quantity(), 10.08), "door areas 1.827 and 10.08 m2")

lines, _ = build(E("Wall linings - Gyprock Plus 10mm Plasterboard-2440mm H-U2-Supply and Install", 128.1, "M"))
check(near(lines[0].quantity(), 312.564) and lines[0].group == "INTERNAL LININGS", "wall lining area = run x 2.44 (312.564 m2)")
lines, _ = build(E("Wall Tiles - 2440mm H - U2", 26.43, "M"))
check(sorted(l.group for l in lines) == ["WALL AND FLOOR TILING", "WATERPROOFING"] and near(lines[0].quantity(), 64.49), "wall tiles give tiling and waterproofing, 64.49 m2 each")
lines, _ = build(E("Soffit Paint Finish - 450mm W-U2", 36.16, "M"))
check(near(lines[0].quantity(), 16.272), "soffit paint = length x 450 mm width (16.272 m2)")
lines, _ = build(E("Internal Wall Paint Finish-U2 = 409m2", 1.0, "EA"), E("Door Paint Finish = 16m2 - U1", 1.0, "EA"))
check([round(l.quantity(), 1) for l in lines] == [409.0, 16.0], "an area written in the label is used when the quantity column is a count of 1")
lines, _ = build(E("90mm Dia Downpipe - U1-3000mm L", 8.0, "EA"), E("Muchugh Jamb Studs - U2", 20.0, "EA"), E("Mchugh Post inside Pier - 2700mm H - U2", 2.0, "EA"))
check(sorted(round(l.quantity(), 2) for l in lines) == [5.4, 24.0, 48.8], "downpipes 8 x 3 m, posts 2 x 2.7 m, jamb studs 20 x 2.44 m are lengths")
lines, _ = build(E("Mchugh Tie-downs @ 600ctrs MAX - U2=190 each", 1.0, "EA"), E("Kitchen Island - 3200 x 900 x 900 - U1", 1.0, "EA"))
check([(l.unit, l.quantity()) for l in lines] == [("each", 190.0), ("each", 1.0)] or sorted((l.unit, l.quantity()) for l in lines) == [("each", 1.0), ("each", 190.0)],
      "a count written in the label (=190 each) and plain joinery counts")
lines, _ = build(E("Mchugh Roof Battens @ 700ctrs MAX-U2", 274.1, "M"))
check(near(one(lines, "ROOF BATTENS").quantity(), 274.1) and near(one(lines, "SCREWS").quantity(), 456.83, 0.1), "roof battens and their screws (1 per 600 mm)")

# rows the rules cannot use are skipped with a reason, never guessed
lines, sk = build(E("DCP & flashings to comply with 5.7.5 and 5.7.4 NCC Vol 2.", 1.0, "EA"),
                  E("Brick Sill - U1", 12.7, "EA"),
                  E("Wall linings - Gyprock Plus 10mm Plasterboard-U1", 120.0, "M"),
                  E("Aluminum Sliding Window/with Mesh - no size", 1.0, "EA"))
check(not lines and len(sk) == 4 and all(why for _d, why in sk), "notes text, a wrong unit, a missing height and a missing size are skipped with a reason")

# the whole Clinton legend
sample = "samples/22l524sp-clinton/ai_entries_zoomed.json"
entries = [e for p in json.load(open(sample)) for e in p["entries"]]
lines, sk = build_takeoff(entries, RES)
check(len(lines) > 200 and len(sk) <= 3, f"the Clinton legend ({len(entries)} rows) gives {len(lines)} lines, {len(sk)} skipped")
bad = [l for l in lines if not near(qty(l.calc, n_of(l)), l.quantity(), 1e-6)]
check(not bad, f"every line's quantity follows from its own inputs through calc.qty ({len(lines) - len(bad)}/{len(lines)})")
check({l.group for l in lines} <= set(RES.groups), "every line is in a residential trade group")
check(all(l.kind in ("res_len", "res_area", "res_count") for l in lines), "every line says how it is measured on the plan")
by_el = {}
for l in lines:
    by_el.setdefault(l.element, set()).add(l.kind)
check(all(len(k) == 1 for k in by_el.values()), "all the lines of one legend row are measured the same way, so one drawn shape feeds them all")

# which profile a legend belongs to: the one whose rules understand more of it
def skipped(entries, name):
    return len(build_takeoff(entries, PROFILES[name])[1])
best = min(PROFILES, key=lambda n: skipped(entries, n))
check(best == "residential", f"the Clinton legend belongs to the residential profile ({ {n: skipped(entries, n) for n in PROFILES} } skipped)")
for sample_dir in ("boxshall-brighton", "82-westgarth"):
    pages = json.load(open(f"samples/{sample_dir}/entries_manual.json"))
    es = [e for p in pages for e in p["entries"]]
    check(min(PROFILES, key=lambda n: skipped(es, n)) != "residential", f"the {sample_dir} footing legend does not go to the residential rules")

# export: the takeoff workbook takes residential groups and lines
path = os.path.join(tempfile.mkdtemp(), "t.xlsx")
write_structural_takeoff(lines, path, sk, RES)
ws = load_workbook(path)["TAKEOFF"]
groups_in_sheet = [c.value for row in ws.iter_rows(min_col=1, max_col=1) for c in row if c.value in RES.groups]
check(set(groups_in_sheet) == {l.group for l in lines}, "the workbook has a section for every residential trade used")
nums = [c.value for row in ws.iter_rows(min_row=2) for c in row[3:10] if isinstance(c.value, (int, float))]
check(len(nums) > len(lines), f"the workbook carries the dimensions ({len(nums)} numbers for {len(lines)} lines)")

print("ALL PASSED" if not fails else f"{fails} FAILED")
raise SystemExit(1 if fails else 0)
