"""Round-trip test: blank out the Westgarth takeoff's dimension cells, refill from our derived lines,
and compare to the estimator's original values."""
import json
import openpyxl

from core.takeoff_builder import PROFILES, build_takeoff
from core.template_filler import fill_template
from core.template_reader import read_template

SRC = "samples/82-westgarth/expected.xlsx"
BLANK = "samples/82-westgarth/template_blank.xlsx"
OUT = "samples/82-westgarth/filled.xlsx"

tpl = read_template(SRC)
sheet = [s for s in tpl.sheets if s.kind == "takeoff"][0]
wb = openpyxl.load_workbook(SRC)
ws = wb[sheet.name]
for it in sheet.items:
    for _, addr in it.input_cells:
        ws[addr].value = None
wb.save(BLANK)

pages = json.load(open("samples/82-westgarth/entries_manual.json"))
lines, _ = build_takeoff([e for p in pages for e in p["entries"]], PROFILES["westgarth"])
rep = fill_template(BLANK, OUT, lines)
print({k: v for k, v in rep.items() if k not in ("skipped", "matches")})

orig = openpyxl.load_workbook(SRC)[sheet.name]
new = openpyxl.load_workbook(OUT)[sheet.name]
ok = bad = 0
bad_ex = []
for m in rep["matches"]:
    for col in ("E", "F", "G", "H", "I"):
        a = f"{col}{m.item.row}"
        o, n = orig[a].value, new[a].value
        if isinstance(o, str) or isinstance(n, str) or (o is None and n is None):
            continue
        if o is not None and n is not None and abs(o - n) <= max(0.06, 0.02 * abs(o)):
            ok += 1
        else:
            bad += 1
            bad_ex.append((a, m.item.description[:30], o, n, m.why))
print("filled cells equal to the original:", ok, "differ:", bad)
for b in bad_ex[:8]:
    print("  ", b)
